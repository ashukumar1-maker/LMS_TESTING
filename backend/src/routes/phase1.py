"""Phase 1 training recommendation API."""

from __future__ import annotations

import json
from dataclasses import dataclass
from http import HTTPStatus
from io import BytesIO
from typing import TYPE_CHECKING, Annotated, cast

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.auth.verify_user import verify_user
from src.get_async_db import get_async_db
from src.models import (
    CourseRepositoryItem,
    Employee,
    EmployeeStenScore,
    Recommendation,
    TnaFormSubmission,
)

if TYPE_CHECKING:
    from collections.abc import Iterable, Iterator

GAP_THRESHOLD = 5
MANDATORY_RECOMMENDATION_COUNT = 5
MIN_TNA_WORD_COUNT = 50
LEVEL_1 = 1
LEVEL_2 = 2
LEVEL_3 = 3
TRAITS: tuple[tuple[str, str], ...] = (
    ("accountability", "Accountability"),
    ("ethical_propensity", "Ethical Propensity"),
    ("self_assurance", "Self-assurance"),
    ("stress_management", "Stress Management"),
    ("learning_agility", "Learning Agility"),
    ("initiative_taking", "Initiative Taking"),
    ("collaboration", "Collaboration"),
    ("inclusion", "Inclusion"),
    ("networking", "Networking"),
    ("interpersonal_skills", "Interpersonal Skills"),
    ("process_orientation", "Process Orientation"),
    ("result_orientation", "Result Orientation"),
    ("planning", "Planning"),
    ("service_orientation", "Service Orientation"),
    ("quality_focus", "Quality Focus"),
    ("flexibility", "Flexibility"),
    ("innovation", "Innovation"),
)
TRAIT_KEYS = tuple(key for key, _ in TRAITS)

DEFINED_SKILLS: list[str] = [
    "Learning Agility",
    "Service Orientation",
    "Process Orientation",
    "Planning",
    "Interpersonal Skills/ Social Competence",
    "Stress Management",
    "Communication",
    "Functional Proficiency",
    "Ethical Propensity",
    "Flexibility",
    "Quality Focus",
    "Training Delivery Skills",
    "Leadership",
    "Coaching skills",
    "Accountability",
    "Self-assurance",
    "Collaboration",
    "Critical thinking",
    "Functional Proficiency - Language",
    "Presentation skills",
    "Strategic Thinking",
    "Technical Skills",
    "Corporate skills",
    "Inclusion",
    "Self Management",
    "Communication Skills",
    "Negotiation Skills",
    "Technical Proficiency",
    "Result Orientation",
    "Networking",
    "Initiative Taking",
    "Innovation",
]

APPROVED_HR_IDS = {"HR-1001", "HR-2024", "PROVANA-HR"}


hr_router = APIRouter(
    prefix="/phase1/hr",
    tags=["phase1-hr"],
    dependencies=[Depends(verify_user)],
)
public_router = APIRouter(prefix="/phase1/public", tags=["phase1-public"])


class UploadSummary(BaseModel):
    """Upload result summary."""

    loaded: int
    invalid: int = 0




class TnaSubmissionCreate(BaseModel):
    """External TNA form submission payload."""

    employee_email: str
    provana_email: str
    employee_name: str = ""
    employee_id_entered: str = ""
    department: str
    team: str
    employee_role: str
    supervisor_name: str
    assessment_period: str
    gap_description: str
    business_impact: str
    capabilities_required: str
    expected_training_outcome: str
    additional_comments: str = ""
    hr_id_entered: str


@dataclass(frozen=True)
class ParsedMettlRecord:
    """Parsed employee and score data from a Mettl workbook."""

    first_name: str
    last_name: str
    email: str
    test_taker_id: str
    overall_status: str
    overall_score: float | None
    test_date: str
    gender: str
    date_of_birth: str
    contact_number: str
    scores: dict[str, float | None]
    strengths_text: str
    areas_of_development_text: str
    response_style: str
    gaps: list[str]


@hr_router.post("/mettl-upload", status_code=HTTPStatus.OK)
async def upload_mettl_file(
    file: Annotated[UploadFile, File(description="Mettl Personality Profiler Excel file")],
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> UploadSummary:
    """Load employees and STEN scores from the Mettl workbook.

    Returns:
        Count of loaded and invalid employee records.
    """
    workbook = _load_xlsx(await file.read())
    records = _parse_mettl_workbook(workbook)
    loaded = 0
    invalid = 0
    for record in records:
        await _upsert_employee_record(db, record)
        loaded += 1
        invalid += 0 if record.overall_status.casefold() == "completed" else 1

    await db.commit()
    await regenerate_recommendations(db)
    return UploadSummary(loaded=loaded, invalid=invalid)


@hr_router.post("/courses-upload", status_code=HTTPStatus.OK)
async def upload_courses_file(
    file: Annotated[UploadFile, File(description="Learning Content Repository Excel file")],
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> UploadSummary:
    """Load courses from the Learning Content Repository workbook.

    Returns:
        Count of loaded course records.
    """
    workbook = _load_xlsx(await file.read())
    courses = _parse_course_workbook(workbook)
    _ = await db.execute(delete(Recommendation))
    _ = await db.execute(delete(CourseRepositoryItem))
    db.add_all(courses)
    await db.commit()
    await regenerate_recommendations(db)
    return UploadSummary(loaded=len(courses))




@hr_router.get("/dashboard", status_code=HTTPStatus.OK)
async def get_dashboard(
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> dict[str, object]:
    """Return employees, courses, recommendations, and TNA submissions for HR."""
    employees = list((await db.execute(select(Employee))).scalars().all())
    scores = {
        score.employee_id: score
        for score in (await db.execute(select(EmployeeStenScore))).scalars().all()
    }
    courses = list((await db.execute(select(CourseRepositoryItem))).scalars().all())
    recommendations = list((await db.execute(select(Recommendation))).scalars().all())
    submissions = list(
        (await db.execute(select(TnaFormSubmission).order_by(TnaFormSubmission.submitted_at.desc())))
        .scalars()
        .all()
    )

    return {
        "traits": [{"key": key, "label": label} for key, label in TRAITS],
        "defined_skills": DEFINED_SKILLS,
        "employees": [_employee_response(employee, scores.get(employee.employee_id)) for employee in employees],
        "courses": [_course_response(course) for course in courses],
        "recommendations": _recommendation_response(recommendations, courses),
        "submissions": [_submission_response(submission) for submission in submissions],
    }


@hr_router.get("/dashboard/summary", status_code=HTTPStatus.OK)
async def get_dashboard_summary(
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> dict[str, object]:
    """Return high-level stats for the dashboard overview cards."""
    from datetime import UTC, datetime, timedelta

    now = datetime.now(UTC).replace(tzinfo=None)
    seven_days_ago = now - timedelta(days=7)

    total_employees = (await db.execute(select(func.count(Employee.employee_id)))).scalar() or 0
    assessed_employees = (
        await db.execute(select(func.count(Employee.employee_id)).where(Employee.is_valid.is_(True)))
    ).scalar() or 0
    new_employees_7d = (
        await db.execute(
            select(func.count(Employee.employee_id)).where(Employee.created_at >= seven_days_ago)
        )
    ).scalar() or 0
    total_courses = (
        await db.execute(select(func.count(CourseRepositoryItem.course_id)))
    ).scalar() or 0
    total_recommendations = (
        await db.execute(select(func.count(Recommendation.recommendation_id)))
    ).scalar() or 0
    total_tna_submissions = (
        await db.execute(select(func.count(TnaFormSubmission.submission_id)))
    ).scalar() or 0
    pending_tna = (
        await db.execute(
            select(func.count(TnaFormSubmission.submission_id)).where(
                TnaFormSubmission.ai_processed.is_(False)
            )
        )
    ).scalar() or 0
    processed_tna = (
        await db.execute(
            select(func.count(TnaFormSubmission.submission_id)).where(
                TnaFormSubmission.ai_processed.is_(True)
            )
        )
    ).scalar() or 0

    # Employees with at least one recommendation
    employees_with_recs = (
        await db.execute(
            select(func.count(func.distinct(Recommendation.employee_email)))
        )
    ).scalar() or 0
    coverage_pct = round((employees_with_recs / assessed_employees * 100) if assessed_employees > 0 else 0, 1)

    return {
        "total_employees": total_employees,
        "assessed_employees": assessed_employees,
        "new_employees_7d": new_employees_7d,
        "total_courses": total_courses,
        "total_recommendations": total_recommendations,
        "total_tna_submissions": total_tna_submissions,
        "pending_tna": pending_tna,
        "processed_tna": processed_tna,
        "employees_with_recommendations": employees_with_recs,
        "recommendation_coverage_pct": coverage_pct,
    }


@hr_router.post("/tna/{submission_id}/process", status_code=HTTPStatus.OK)
async def process_tna_ai(
    submission_id: str,
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> dict[str, object]:
    """Re-run AI trait extraction on a TNA submission."""
    from src.ai_processing import process_tna_submission

    submission = await db.get(TnaFormSubmission, submission_id)
    if submission is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Submission not found.")

    traits = await process_tna_submission(db, submission_id)
    await regenerate_recommendations(db)
    return {"status": "processed", "extracted_traits": traits}


@public_router.get("/employee-lookup", status_code=HTTPStatus.OK)
async def lookup_employee(
    query: Annotated[str, Query(min_length=1)],
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> dict[str, str]:
    """Lookup one employee by email address or Mettl test-taker ID.

    Returns:
        Employee name and email for the external form.

    Raises:
        HTTPException: If the employee is not found.
    """
    normalized_query = query.strip().casefold()
    result = await db.execute(
        select(Employee).where(
            (Employee.email == normalized_query)
            | (Employee.test_taker_id == query.strip())
            | (Employee.employee_id == query.strip())
        )
    )
    employee = result.scalar_one_or_none()
    if employee is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Employee not found.")

    return {
        "employee_id": employee.employee_id,
        "employee_email": employee.email,
        "employee_name": _full_name(employee.first_name, employee.last_name),
    }


@public_router.post("/tna-submissions", status_code=HTTPStatus.CREATED)
async def submit_tna_form(
    submission: TnaSubmissionCreate,
    db: Annotated[AsyncSession, Depends(get_async_db)],
) -> dict[str, str]:
    """Save a verified external Training Needs Assessment form submission.

    Returns:
        Save status payload.

    Raises:
        HTTPException: If validation, employee lookup, or HR verification fails.
    """
    _validate_tna_submission(submission)

    # Validate provana email
    if not submission.provana_email or not submission.provana_email.strip().lower().endswith("@provana.com"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Provana Email is required and must be a valid @provana.com email address.",
        )

    # Match by email first, then by employee_id
    employee = (
        await db.execute(
            select(Employee).where(Employee.email == submission.employee_email.casefold())
        )
    ).scalar_one_or_none()
    if employee is None and submission.employee_id_entered:
        employee = (
            await db.execute(
                select(Employee).where(Employee.employee_id == submission.employee_id_entered.strip())
            )
        ).scalar_one_or_none()
    if employee is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Employee not found.")

    is_verified = submission.hr_id_entered in APPROVED_HR_IDS
    if not is_verified:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="HR verification failed. Enter an approved HR ID.",
        )

    tna_record = TnaFormSubmission(
        employee_email=employee.email,
        provana_email=submission.provana_email.strip().lower(),
        employee_name=submission.employee_name or _full_name(employee.first_name, employee.last_name),
        employee_id_entered=submission.employee_id_entered,
        department=submission.department,
        team=submission.team,
        employee_role=submission.employee_role,
        supervisor_name=submission.supervisor_name,
        assessment_period=submission.assessment_period,
        gap_description=submission.gap_description,
        business_impact=submission.business_impact,
        capabilities_required=submission.capabilities_required,
        expected_training_outcome=submission.expected_training_outcome,
        additional_comments=submission.additional_comments,
        hr_id_entered=submission.hr_id_entered,
        is_verified=True,
    )
    db.add(tna_record)
    await db.flush()

    # Trigger AI trait extraction
    from src.ai_processing import process_tna_submission

    await process_tna_submission(db, tna_record.submission_id)
    await regenerate_recommendations(db)
    return {"status": "saved"}


async def regenerate_recommendations(db: AsyncSession) -> None:
    """Regenerate recommendations based on TNA skill extraction + Mettl gaps."""
    _ = await db.execute(delete(Recommendation))
    employees = list(
        (await db.execute(select(Employee).where(Employee.is_valid.is_(True)))).scalars().all()
    )
    scores_by_employee_id = {
        score.employee_id: score
        for score in (await db.execute(select(EmployeeStenScore))).scalars().all()
    }
    courses = list((await db.execute(select(CourseRepositoryItem))).scalars().all())
    submissions = {
        submission.employee_email: submission
        for submission in (
            await db.execute(
                select(TnaFormSubmission)
                .where(TnaFormSubmission.is_verified.is_(True))
                .order_by(TnaFormSubmission.submitted_at.desc())
            )
        )
        .scalars()
        .all()
    }

    for employee in employees:
        score = scores_by_employee_id.get(employee.employee_id)
        submission = submissions.get(employee.email)

        # No TNA submission = no department info = cannot filter courses properly
        if submission is None:
            continue

        # Get skills from TNA AI extraction
        extracted_skills: list[str] = []
        if submission.extracted_skills:
            extracted_skills = _json_list(submission.extracted_skills)

        # Get Mettl trait gaps
        mettl_gaps: list[str] = []
        if score:
            mettl_gaps = _json_list(score.gaps_json)

        # Skill-based recommendations from TNA
        if extracted_skills:
            matches = _rank_skill_matches(extracted_skills, courses, submission)
            for index, match in enumerate(matches, start=1):
                db.add(
                    Recommendation(
                        employee_email=employee.email,
                        course_id=match.course.course_id,
                        recommendation_source="tna",
                        matched_skills=json.dumps(match.matched_skills),
                        matched_traits=json.dumps([]),
                        recommendation_rank=index,
                        is_mandatory=index <= MANDATORY_RECOMMENDATION_COUNT,
                    )
                )
        elif mettl_gaps:
            # Fallback: use Mettl trait gaps mapped to skills, still filtered by department
            matches = _rank_skill_matches(
                [_trait_to_skill(g) for g in mettl_gaps], courses, submission
            )
            for index, match in enumerate(matches, start=1):
                db.add(
                    Recommendation(
                        employee_email=employee.email,
                        course_id=match.course.course_id,
                        recommendation_source="mettl",
                        matched_skills=json.dumps(match.matched_skills),
                        matched_traits=json.dumps(mettl_gaps),
                        recommendation_rank=index,
                        is_mandatory=index <= MANDATORY_RECOMMENDATION_COUNT,
                    )
                )

    await db.commit()


@dataclass(frozen=True)
class CourseMatch:
    """Course match with matched skills and score."""

    course: CourseRepositoryItem
    matched_skills: list[str]
    score: float = 0.0


def _load_xlsx(content: bytes) -> Workbook:
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty.")
    try:
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file could not be parsed as an .xlsx workbook.",
        ) from exc


def _parse_mettl_workbook(workbook: Workbook) -> list[ParsedMettlRecord]:
    if "Test Takers" not in workbook.sheetnames or "Analysis" not in workbook.sheetnames:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mettl workbook must include Test Takers and Analysis sheets.",
        )
    test_taker_rows = _rows_by_email(workbook["Test Takers"], header_row_number=4)
    analysis_rows = _rows_by_email(workbook["Analysis"], header_row_number=4)
    trait_columns = _analysis_trait_columns(workbook["Analysis"])

    records: list[ParsedMettlRecord] = []
    for email, row in test_taker_rows.items():
        analysis = analysis_rows.get(email, ())
        scores = {
            trait_key: _number_or_none(_value_at(analysis, column_index))
            for trait_key, column_index in trait_columns.items()
        }
        gaps = [
            trait_key
            for trait_key, score in scores.items()
            if score is not None and score < GAP_THRESHOLD
        ]
        records.append(
            ParsedMettlRecord(
                first_name=_value_at(row, 0),
                last_name=_value_at(row, 1),
                email=email,
                test_taker_id=_value_at(row, 3),
                overall_status=_value_at(row, 4),
                overall_score=_number_or_none(_value_at(row, 6)),
                test_date=_value_at(row, 9),
                gender=_value_at(row, 20),
                date_of_birth=_value_at(row, 18),
                contact_number=_value_at(row, 19),
                scores=scores,
                strengths_text="\n".join(_non_empty_texts(_value_at(analysis, i) for i in range(5, 8))),
                areas_of_development_text="\n".join(
                    _non_empty_texts(_value_at(analysis, i) for i in range(8, 11))
                ),
                response_style=_value_at(analysis, 3),
                gaps=gaps,
            )
        )

    return records


def _parse_course_workbook(workbook: Workbook) -> list[CourseRepositoryItem]:
    # Support both old and new sheet names
    sheet_name = None
    for name in workbook.sheetnames:
        if "content repository" in name.casefold():
            sheet_name = name
            break
    worksheet = cast(
        "Worksheet | None",
        workbook[sheet_name] if sheet_name else workbook.active,
    )
    if worksheet is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Course workbook does not contain a worksheet.",
        )
    rows = cast("Iterable[tuple[object, ...]]", worksheet.iter_rows(min_row=2, values_only=True))
    courses: list[CourseRepositoryItem] = []
    for row in rows:
        course_name = _value_at(row, 0)
        if not course_name:
            continue
        courses.append(
            CourseRepositoryItem(
                course_name=course_name,
                description=_value_at(row, 1),
                department=_value_at(row, 2),
                role_type=_value_at(row, 3),
                level_1=_yes(_value_at(row, 4)),
                level_2=_yes(_value_at(row, 5)),
                level_3=_yes(_value_at(row, 6)),
                content_format=_value_at(row, 7),
                is_lms_module=_yes(_value_at(row, 8)),
                training_duration_minutes=_integer_or_none(_value_at(row, 9)),
                primary_skill=_value_at(row, 10),
                secondary_skill=_value_at(row, 11),
            )
        )
    return courses


async def _upsert_employee_record(db: AsyncSession, record: ParsedMettlRecord) -> None:
    employee = (
        await db.execute(select(Employee).where(Employee.email == record.email))
    ).scalar_one_or_none()
    if employee is None:
        employee = Employee(email=record.email)
        db.add(employee)
        await db.flush()

    employee.first_name = record.first_name
    employee.last_name = record.last_name
    employee.test_taker_id = record.test_taker_id
    employee.overall_status = record.overall_status
    employee.overall_score = record.overall_score
    employee.test_date = record.test_date
    employee.gender = record.gender
    employee.date_of_birth = record.date_of_birth
    employee.contact_number = record.contact_number
    employee.is_valid = record.overall_status.casefold() == "completed"

    sten_score = await db.get(EmployeeStenScore, employee.employee_id)
    if sten_score is None:
        sten_score = EmployeeStenScore(employee_id=employee.employee_id)
        db.add(sten_score)
    for trait_key in TRAIT_KEYS:
        setattr(sten_score, trait_key, record.scores.get(trait_key))
    sten_score.strengths_text = record.strengths_text
    sten_score.areas_of_development_text = record.areas_of_development_text
    sten_score.response_style = record.response_style
    sten_score.gaps_json = json.dumps(record.gaps)


def _rank_skill_matches(
    skills: list[str],
    courses: list[CourseRepositoryItem],
    submission: TnaFormSubmission | None,
) -> list[CourseMatch]:
    matches: list[CourseMatch] = []
    normalized_skills = [s.casefold().strip() for s in skills]

    for course in courses:
        if not _course_applies_to_employee(course, submission):
            continue

        score = 0.0
        matched: list[str] = []
        primary = course.primary_skill.casefold().strip()
        secondary = course.secondary_skill.casefold().strip()

        for skill in normalized_skills:
            if _skill_matches(skill, primary):
                score += 2.0
                matched.append(course.primary_skill)
            elif _skill_matches(skill, secondary):
                score += 1.0
                matched.append(course.secondary_skill)

        if matched:
            matches.append(CourseMatch(course=course, matched_skills=matched, score=score))

    return sorted(matches, key=lambda m: (-m.score, m.course.course_name))


def _skill_matches(skill: str, course_skill: str) -> bool:
    if not course_skill:
        return False
    if skill == course_skill:
        return True
    if skill in course_skill or course_skill in skill:
        return True
    return False


def _course_applies_to_employee(
    course: CourseRepositoryItem,
    submission: TnaFormSubmission | None,
) -> bool:
    if submission is None:
        return True

    # Department filter: match employee dept or "All Departments"
    if course.department:
        course_dept = course.department.casefold().strip()
        if course_dept != "all departments":
            if submission.department:
                emp_dept = submission.department.casefold().strip()
                # Support comma-separated multi-department courses
                course_depts = [d.strip() for d in course_dept.split(",")]
                if not any(emp_dept == d or emp_dept in d for d in course_depts):
                    return False

    # Level filter
    level = _role_level(submission.employee_role)
    if level == LEVEL_1:
        return course.level_1
    if level == LEVEL_2:
        return course.level_2
    if level == LEVEL_3:
        return course.level_3
    return True


def _trait_to_skill(trait_key: str) -> str:
    """Map a Mettl trait key to its equivalent skill name."""
    mapping = {
        "accountability": "Accountability",
        "ethical_propensity": "Ethical Propensity",
        "self_assurance": "Self-assurance",
        "stress_management": "Stress Management",
        "learning_agility": "Learning Agility",
        "initiative_taking": "Initiative Taking",
        "collaboration": "Collaboration",
        "inclusion": "Inclusion",
        "networking": "Networking",
        "interpersonal_skills": "Interpersonal Skills/ Social Competence",
        "process_orientation": "Process Orientation",
        "result_orientation": "Result Orientation",
        "planning": "Planning",
        "service_orientation": "Service Orientation",
        "quality_focus": "Quality Focus",
        "flexibility": "Flexibility",
        "innovation": "Innovation",
    }
    return mapping.get(trait_key, trait_key)


def _role_level(role: str) -> int | None:
    normalized_role = role.casefold()
    if any(token in normalized_role for token in ["deputy manager", "dm"]):
        return LEVEL_3
    if any(token in normalized_role for token in ["team lead", "assistant manager", "trainer", "qa"]):
        return LEVEL_2
    if normalized_role:
        return LEVEL_1
    return None


def _rows_by_email(worksheet: Worksheet, *, header_row_number: int) -> dict[str, tuple[object, ...]]:
    rows: dict[str, tuple[object, ...]] = {}
    for row in cast(
        "Iterable[tuple[object, ...]]",
        worksheet.iter_rows(min_row=header_row_number + 1, values_only=True),
    ):
        email = _normalize_email(_value_at(row, 2))
        if email:
            rows[email] = row
    return rows


def _analysis_trait_columns(worksheet: Worksheet) -> dict[str, int]:
    row = next(
        cast("Iterator[tuple[object, ...]]", worksheet.iter_rows(min_row=5, max_row=5, values_only=True)),
        (),
    )
    columns: dict[str, int] = {}
    for index, _value in enumerate(row):
        trait_key = _trait_key_for_label(_value_at(row, index))
        if trait_key:
            columns[trait_key] = index
    return columns


def _trait_key_for_label(label: str) -> str | None:
    normalized = label.casefold().replace("/", " ")
    if "interpersonal skills" in normalized:
        return "interpersonal_skills"
    for trait_key, trait_label in TRAITS:
        if trait_label.casefold() in normalized:
            return trait_key
    return None


def _employee_response(employee: Employee, score: EmployeeStenScore | None) -> dict[str, object]:
    return {
        "employee_id": employee.employee_id,
        "first_name": employee.first_name,
        "last_name": employee.last_name,
        "full_name": _full_name(employee.first_name, employee.last_name),
        "email": employee.email,
        "test_taker_id": employee.test_taker_id,
        "overall_status": employee.overall_status,
        "overall_score": employee.overall_score,
        "test_date": employee.test_date,
        "gender": employee.gender,
        "date_of_birth": employee.date_of_birth,
        "contact_number": employee.contact_number,
        "is_valid": employee.is_valid,
        "scores": _score_dict(score),
        "gap_traits": _json_list(score.gaps_json if score else "[]"),
        "strengths_text": score.strengths_text if score else "",
        "areas_of_development_text": score.areas_of_development_text if score else "",
        "response_style": score.response_style if score else "",
    }


def _score_dict(score: EmployeeStenScore | None) -> dict[str, float | None]:
    if score is None:
        return {}
    return {trait: getattr(score, trait) for trait in TRAIT_KEYS}


def _course_response(course: CourseRepositoryItem) -> dict[str, object]:
    return {
        "course_id": course.course_id,
        "course_name": course.course_name,
        "description": course.description,
        "department": course.department,
        "role_type": course.role_type,
        "level_1": course.level_1,
        "level_2": course.level_2,
        "level_3": course.level_3,
        "content_format": course.content_format,
        "is_lms_module": course.is_lms_module,
        "training_duration_minutes": course.training_duration_minutes,
        "primary_skill": course.primary_skill,
        "secondary_skill": course.secondary_skill,
    }


def _recommendation_response(
    recommendations: list[Recommendation],
    courses: list[CourseRepositoryItem],
) -> dict[str, list[dict[str, object]]]:
    course_by_id = {course.course_id: course for course in courses}
    grouped: dict[str, list[dict[str, object]]] = {}
    for recommendation in sorted(recommendations, key=lambda item: item.recommendation_rank):
        course = course_by_id.get(recommendation.course_id)
        if course is None:
            continue
        grouped.setdefault(recommendation.employee_email, []).append(
            {
                "recommendation_id": recommendation.recommendation_id,
                "course": _course_response(course),
                "source": recommendation.recommendation_source,
                "matched_skills": _json_list(recommendation.matched_skills),
                "matched_traits": _json_list(recommendation.matched_traits),
                "rank": recommendation.recommendation_rank,
                "is_mandatory": recommendation.is_mandatory,
            }
        )
    return grouped


def _submission_response(submission: TnaFormSubmission) -> dict[str, object]:
    return {
        "submission_id": submission.submission_id,
        "employee_email": submission.employee_email,
        "provana_email": submission.provana_email,
        "employee_name": submission.employee_name,
        "department": submission.department,
        "team": submission.team,
        "employee_role": submission.employee_role,
        "supervisor_name": submission.supervisor_name,
        "assessment_period": submission.assessment_period,
        "gap_description": submission.gap_description,
        "business_impact": submission.business_impact,
        "capabilities_required": submission.capabilities_required,
        "expected_training_outcome": submission.expected_training_outcome,
        "additional_comments": submission.additional_comments,
        "hr_id_entered": submission.hr_id_entered,
        "is_verified": submission.is_verified,
        "submitted_at": submission.submitted_at.isoformat(),
        "ai_processed": submission.ai_processed,
        "extracted_gaps": _json_list(submission.extracted_traits or "[]"),
        "extracted_skills": _json_list(submission.extracted_skills or "[]"),
        "unmapped_skills": _json_list(submission.unmapped_skills or "[]"),
    }


def _validate_tna_submission(submission: TnaSubmissionCreate) -> None:
    required_values = (
        submission.gap_description,
        submission.business_impact,
        submission.capabilities_required,
        submission.expected_training_outcome,
    )
    for value in required_values:
        if len(value.split()) < MIN_TNA_WORD_COUNT:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Each required narrative response must contain at least 50 words.",
            )


def _value_at(row_values: tuple[object, ...], index: int) -> str:
    if index < 0 or index >= len(row_values):
        return ""
    value = row_values[index]
    if value is None:
        return ""
    return str(value).strip()


def _normalize_email(value: str) -> str:
    normalized = value.casefold()
    if "@" not in normalized or "." not in normalized:
        return ""
    return normalized


def _number_or_none(value: str) -> float | None:
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _integer_or_none(value: str) -> int | None:
    number = _number_or_none(value)
    if number is None:
        return None
    return int(number)


def _yes(value: str) -> bool:
    return value.casefold() in {"yes", "y", "true", "1"}


def _non_empty_texts(values: Iterable[str]) -> list[str]:
    return [value for value in values if value and value.upper() != "NONE"]


def _json_list(value: str) -> list[str]:
    try:
        parsed = cast("object", json.loads(value))
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    parsed_items = cast("list[object]", parsed)
    return [item for item in parsed_items if isinstance(item, str)]




def _full_name(first_name: str, last_name: str) -> str:
    return " ".join([part for part in [first_name, last_name] if part])


router = APIRouter()
router.include_router(hr_router)
router.include_router(public_router)
